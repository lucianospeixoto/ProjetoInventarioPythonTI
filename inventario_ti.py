import os
import platform
import socket
import psutil
import wmi
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

# Caminho da planilha na pasta local do projeto
PLANILHA_PATH = r'C:\Users\luciano.peixoto\Documents\inventario_python\inventario.xlsx'

COLUNAS = [
    'Hostname', 'IP', 'Domínio/Workgroup',
    'Sistema', 'Versão', 'Placa-mãe', 'Serial BIOS',
    'Processador', 'Núcleos Físicos', 'Núcleos Lógicos', 'RAM Total (GB)',
    'Disco 1', 'Capacidade Disco 1 (GB)', 'Disco 2', 'Capacidade Disco 2 (GB)',
    'Espaço Total C: (GB)', 'Espaço Livre C: (GB)',
    'Placa de Vídeo 1', 'Placa de Vídeo 2',
]
LARGURAS = [20, 15, 18, 25, 15, 25, 20, 40, 14, 14, 15, 10, 22, 10, 22, 18, 18, 30, 30]


def criar_ou_abrir_planilha(caminho, colunas, larguras):
    """
    Cria a planilha (com cabeçalho e larguras de coluna) se ela ainda não
    existir na pasta local, ou abre a planilha existente.
    """
    pasta = os.path.dirname(caminho)
    if pasta and not os.path.exists(pasta):
        os.makedirs(pasta, exist_ok=True)

    if os.path.exists(caminho):
        wb = load_workbook(caminho)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = 'Inventário'
        ws.append(colunas)

    for i, largura in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(i)].width = largura

    return wb, ws


def coletar_inventario():
    dados = {}

    print(' ')
    print('================== ESPECIFICAÇÕES DO SISTEMA ================== ')
    print(' ')

    # ================== IDENTIFICAÇÃO DA MÁQUINA ==================
    print(' ')
    print('================== IDENTIFICAÇÃO ================== ')
    print(' ')

    hostname = socket.gethostname()
    try:
        ip_local = socket.gethostbyname(hostname)
    except socket.gaierror:
        ip_local = 'Não identificado'

    print('Hostname: ', hostname)
    print('IP: ', ip_local)

    dados['hostname'] = hostname
    dados['ip'] = ip_local

    # ================== SISTEMA ==================
    print(' ')
    print('================== SISTEMA ================== ')
    print(' ')

    so = wmi.WMI()
    os_name = ''
    os_version = ''
    dominio_workgroup = 'Não identificado'
    for os_info in so.Win32_OperatingSystem():
        os_name = os_info.Caption
        os_version = os_info.Version
        print('Nome: ', os_info.Caption)
        print('Versão: ', os_info.Version)

    for cs in so.Win32_ComputerSystem():
        dominio_workgroup = cs.Domain
        print('Domínio/Workgroup: ', cs.Domain)

    dados['os_name'] = os_name
    dados['os_version'] = os_version
    dados['dominio_workgroup'] = dominio_workgroup

    # ================== BIOS / PLACA-MÃE ==================
    print(' ')
    print('================== BIOS / PLACA-MÃE ================== ')
    print(' ')

    serial_bios = 'Não identificado'
    for bios in so.Win32_BIOS():
        serial_bios = bios.SerialNumber
        print('Número de série (BIOS): ', bios.SerialNumber)
        print('Versão BIOS: ', bios.SMBIOSBIOSVersion)

    placa_mae = fabricante_placa = ''
    for board in so.Win32_BaseBoard():
        placa_mae = board.Product
        fabricante_placa = board.Manufacturer
        print('Placa-mãe: ', board.Manufacturer, board.Product)

    dados['serial_bios'] = serial_bios
    dados['placa_mae'] = f'{fabricante_placa} {placa_mae}'.strip()

    # ================== PROCESSADOR ==================
    print(' ')
    print('================== PROCESSADOR ================== ')
    print(' ')

    p = wmi.WMI()
    cpu_name = ''
    for cpu in p.Win32_Processor():
        cpu_name = cpu.Name
        print('Processador: ', cpu.Name)
    print('Núcleos físicos: ', psutil.cpu_count(logical=False))
    print('Núcleos lógicos: ', psutil.cpu_count(logical=True))

    dados['cpu_name'] = cpu_name
    dados['nucleos_fisicos'] = psutil.cpu_count(logical=False)
    dados['nucleos_logicos'] = psutil.cpu_count(logical=True)

    # ================== MEMÓRIA RAM ==================
    m = wmi.WMI()
    print(' ')
    print('================== MEMÓRIA RAM ================== ')
    print(' ')

    for mem in m.Win32_PhysicalMemory():
        if mem.SMBIOSMemoryType == 24:
            tipo_memoria = 'DDR3'
        elif mem.SMBIOSMemoryType == 26:
            tipo_memoria = 'DDR4'
        elif mem.SMBIOSMemoryType == 34:
            tipo_memoria = 'DDR5'
        else:
            tipo_memoria = 'Desconhecido'
        print('Slot: ', mem.DeviceLocator)
        print('Tipo ', tipo_memoria)
        print('Capacidade: ', round(int(mem.Capacity) / 1073741824, 2), 'GB')
        print(' ')

    ram_total = round(psutil.virtual_memory().total / 1073741824, 2)
    print('RAM Total: ', ram_total, 'GB')
    dados['ram_total'] = ram_total

    # ================== DISCO ==================
    c = wmi.WMI(namespace=r'root\Microsoft\Windows\Storage')
    print(' ')
    print('================== DISCO ================== ')
    print(' ')

    discos = []
    for disk in c.MSFT_PhysicalDisk():
        if disk.MediaType == 4:
            tipo = 'SSD'
        elif disk.MediaType == 3:
            tipo = 'HD'
        else:
            tipo = 'Desconhecido'
        tamanho = round(int(disk.Size) / 1073741824, 2)
        discos.append((tipo, tamanho))
        print('Tipo:', tipo)
        print('Tamanho: ', tamanho, 'GB')
        print(' ')

    disco1_tipo = disco1_tam = ''
    disco2_tipo = disco2_tam = ''
    if len(discos) > 0:
        disco1_tipo, disco1_tam = discos[0]
    if len(discos) > 1:
        disco2_tipo, disco2_tam = discos[1]

    dados['disco1_tipo'] = disco1_tipo
    dados['disco1_tam'] = disco1_tam
    dados['disco2_tipo'] = disco2_tipo
    dados['disco2_tam'] = disco2_tam

    # Espaço livre do disco C:
    try:
        uso_disco = psutil.disk_usage('C:\\')
        espaco_total = round(uso_disco.total / 1073741824, 2)
        espaco_livre = round(uso_disco.free / 1073741824, 2)
    except Exception:
        espaco_total = espaco_livre = 0
    print('Espaço total (C:): ', espaco_total, 'GB')
    print('Espaço livre (C:): ', espaco_livre, 'GB')

    dados['espaco_total'] = espaco_total
    dados['espaco_livre'] = espaco_livre

    # ================== PLACA DE VÍDEO ==================
    print(' ')
    print('================== PLACA DE VÍDEO ================== ')
    print(' ')

    placas_video = []
    for gpu in so.Win32_VideoController():
        placas_video.append(gpu.Name)
        print('Placa de vídeo: ', gpu.Name)

    dados['gpu_1'] = placas_video[0] if len(placas_video) > 0 else ''
    dados['gpu_2'] = placas_video[1] if len(placas_video) > 1 else ''

    print(' ')
    print('Coleta Concluída!')

    return dados


def salvar_no_inventario(dados):
    wb, ws = criar_ou_abrir_planilha(PLANILHA_PATH, COLUNAS, LARGURAS)

    # Se já existe um registro para esse hostname, atualiza (substitui) a linha
    linha_existente = None
    for idx, linha in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if linha[0] == dados['hostname']:
            linha_existente = idx
            break

    if linha_existente:
        ws.delete_rows(linha_existente)
        print('Registro anterior deste hostname foi atualizado.')

    ws.append([
        dados['hostname'], dados['ip'], dados['dominio_workgroup'],
        dados['os_name'], dados['os_version'], dados['placa_mae'], dados['serial_bios'],
        dados['cpu_name'], dados['nucleos_fisicos'], dados['nucleos_logicos'], dados['ram_total'],
        dados['disco1_tipo'], dados['disco1_tam'], dados['disco2_tipo'], dados['disco2_tam'],
        dados['espaco_total'], dados['espaco_livre'],
        dados['gpu_1'], dados['gpu_2'],
    ])

    wb.save(PLANILHA_PATH)
    print('Informações salvas na planilha local com sucesso!')


if __name__ == '__main__':
    dados = coletar_inventario()
    salvar_no_inventario(dados)
    input('Pressione ENTER para fechar o programa...')